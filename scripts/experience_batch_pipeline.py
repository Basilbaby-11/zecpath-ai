"""Batch experience extraction with job relevance and role-to-role similarity.

Run from the repository root after placing the parser and engine in the project:
  python scripts/experience_batch_pipeline.py --input-dir data/sample_resumes_with_Experience

The output keeps the existing all_candidate_experiences.json structure and adds
job_relevance_analysis per candidate plus a top-level role_to_role_similarity list.
"""
from __future__ import annotations
import argparse
import json
import re
import sys
from itertools import combinations
from pathlib import Path
from datetime import date
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from parsers.experience_parser import parse_experience  # noqa: E402
from ats_engine.experience_relevance_engine import (  # noqa: E402
    score_experience_relevance,
    role_similarity,
)

INPUT_DIR = ROOT / "data" / "sample_resumes_with_Experience"
OUTPUT_DIR = ROOT / "data" / "experience_outputs"
SUPPORTED = {".pdf", ".docx", ".txt"}
PREFERENCE = {".docx": 0, ".pdf": 1, ".txt": 2}

# The same two example job profiles used in the previous relevance output.
# Replace/edit these or pass --job-description to score against a specific job.
DEFAULT_JOB_PROFILES = [
    {
        "job_id": "JD_ML_ENGINEER",
        "title": "Machine Learning Engineer",
        "required_skills": ["Python", "Machine Learning", "scikit-learn", "Model evaluation", "Classification"],
        "responsibilities": "Prepare and preprocess datasets, train and evaluate classification models, compare algorithms, track precision recall and F1 score, build reproducible machine learning pipelines, and document experiments.",
        "description": "Entry-level machine learning role requiring practical Python programming, data preprocessing, model training, and evaluation.",
    },
    {
        "job_id": "JD_DATA_ANALYST",
        "title": "Data Analyst",
        "required_skills": ["SQL", "Power BI", "Data cleaning", "Data visualization", "Python"],
        "responsibilities": "Clean and validate datasets, write SQL queries, build dashboards, analyze trends, communicate findings, and document data quality limitations.",
        "description": "Entry-level data analysis role focused on reporting, dashboards, data quality, and business insights.",
    },
]

# Small transparent vocabulary used only to improve scoring when parser evidence
# did not populate skills_used. It does not alter the original parsed field.
SKILL_VOCAB = [
    "python", "java", "javascript", "typescript", "sql", "postgresql", "mysql", "redis",
    "fastapi", "django", "rest api", "spring boot", "react", "scikit-learn", "pytorch",
    "machine learning", "classification", "regression", "model evaluation", "data preprocessing",
    "pandas", "numpy", "nlp", "spacy", "nltk", "tf-idf", "embeddings", "mlflow",
    "power bi", "excel", "data visualization", "data cleaning", "siem", "splunk",
    "incident response", "vulnerability management", "linux", "docker", "aws", "azure",
    "kubernetes", "terraform", "jenkins", "github actions", "prometheus", "grafana", "ci/cd",
    "autocad", "site supervision", "quantity estimation", "boq", "qa/qc", "construction safety",
    "agile", "scrum", "jira", "stakeholder management", "project planning", "risk tracking",
    "requirements analysis", "applicant tracking systems", "recruitment coordination",
    "interview scheduling", "candidate communication", "reporting", "data validation", "hr operations",
    "pytest", "git", "api testing", "hibernate", "junit", "mockito", "maven", "bash", "cloudwatch",
]


def extract_text(path: Path) -> str:
    if path.suffix.lower() == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")
    if path.suffix.lower() == ".pdf":
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages).strip()
    if path.suffix.lower() == ".docx":
        from docx import Document
        doc = Document(path)
        chunks = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                chunks.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
        return "\n".join(chunks).strip()
    raise ValueError(f"Unsupported file type: {path.suffix}")


def safe_stem(path: Path) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", path.stem).strip("_").upper() or "RESUME"


def choose_one_file_per_resume(files: list[Path]) -> tuple[list[Path], list[str]]:
    grouped: dict[str, list[Path]] = {}
    for path in files:
        grouped.setdefault(path.stem.casefold(), []).append(path)
    chosen, skipped = [], []
    for variants in grouped.values():
        variants.sort(key=lambda p: (PREFERENCE.get(p.suffix.lower(), 99), p.name.casefold()))
        chosen.append(variants[0])
        skipped.extend(p.name for p in variants[1:])
    chosen.sort(key=lambda p: p.name.casefold())
    return chosen, sorted(skipped)


def inferred_skills(experience: dict[str, Any]) -> list[str]:
    evidence = " ".join(str(experience.get(k) or "") for k in ("designation", "description", "source_text")).casefold()
    found = []
    for skill in SKILL_VOCAB:
        pattern = r"(?<![a-z0-9+#])" + re.escape(skill.casefold()) + r"(?![a-z0-9+#])"
        if re.search(pattern, evidence, re.I):
            found.append(skill)
    # Use parsed skills first, then transparent evidence-derived additions.
    return list(dict.fromkeys([*(experience.get("skills_used") or []), *found]))


def _role_skill_view(exp: dict[str, Any]) -> dict[str, Any]:
    view = dict(exp)
    view["skills_used"] = inferred_skills(exp)
    return view


def build_role_similarity(candidates: list[dict[str, Any]]) -> list[dict[str, Any]]:
    roles = []
    for candidate in candidates:
        for exp in candidate.get("experiences", []):
            roles.append((candidate, exp, _role_skill_view(exp)))
    comparisons = []
    for (source_candidate, source_exp, source_view), (target_candidate, target_exp, target_view) in combinations(roles, 2):
        forward = role_similarity(source_view.get("designation"), target_view.get("designation"), source_view["skills_used"], target_view["skills_used"])
        reverse = role_similarity(target_view.get("designation"), source_view.get("designation"), target_view["skills_used"], source_view["skills_used"])
        score = round((forward.get("score", 0.0) + reverse.get("score", 0.0)) / 2, 2)
        title_score = round((forward.get("title_similarity", 0.0) + reverse.get("title_similarity", 0.0)) / 2, 2)
        skill_score = round((forward.get("skill_similarity", 0.0) + reverse.get("skill_similarity", 0.0)) / 2, 2)
        source_skills, target_skills = set(source_view["skills_used"]), set(target_view["skills_used"])
        comparisons.append({
            "source_candidate_id": source_candidate["candidate_id"],
            "source_experience_id": source_exp.get("experience_id"),
            "source_role": source_exp.get("designation"),
            "source_employer": source_exp.get("employer"),
            "target_candidate_id": target_candidate["candidate_id"],
            "target_experience_id": target_exp.get("experience_id"),
            "target_role": target_exp.get("designation"),
            "target_employer": target_exp.get("employer"),
            "score": score,
            "relationship": "high" if score >= 65 else "moderate" if score >= 30 else "low",
            "title_similarity": title_score,
            "skill_similarity": skill_score,
            "matched_skills": sorted(source_skills & target_skills),
            "source_only_skills": sorted(source_skills - target_skills),
            "target_only_skills": sorted(target_skills - source_skills),
            "method": "symmetric average of TF-IDF title similarity and normalized skill overlap",
        })
    comparisons.sort(key=lambda row: row["score"], reverse=True)
    return comparisons


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=INPUT_DIR)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    parser.add_argument("--job-description", type=Path, default=None,
                        help="Optional JSON object with title, required_skills, responsibilities")
    args = parser.parse_args()
    args.input_dir.mkdir(parents=True, exist_ok=True)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    all_files = sorted((p for p in args.input_dir.iterdir() if p.is_file() and p.suffix.lower() in SUPPORTED), key=lambda p: p.name.casefold())
    if not all_files:
        print(f"No PDF/DOCX/TXT resumes found in: {args.input_dir}")
        return 2
    files, skipped = choose_one_file_per_resume(all_files)
    if args.job_description:
        job = json.loads(args.job_description.read_text(encoding="utf-8"))
        if not isinstance(job, dict):
            raise ValueError("Job description JSON must contain an object.")
        job_profiles = [job]
    else:
        job_profiles = DEFAULT_JOB_PROFILES

    results, failures = [], []
    for index, path in enumerate(files, start=1):
        candidate_id = f"CAND_{index:03d}"
        try:
            text = extract_text(path)
            if not text.strip():
                raise ValueError("No text extracted; scanned PDFs may require OCR.")
            parsed = parse_experience(text, candidate_id=candidate_id)
            parsed["source_file"] = path.name
            parsed["processing_status"] = "success"
            scoring_experiences = [_role_skill_view(exp) for exp in parsed["experiences"]]
            parsed["job_relevance_analysis"] = [
                {"job_id": profile.get("job_id", f"JD_{idx:03d}"), **score_experience_relevance(scoring_experiences, profile)}
                for idx, profile in enumerate(job_profiles, start=1)
            ]
            results.append(parsed)
            (args.output_dir / f"{candidate_id}_{safe_stem(path)}.json").write_text(json.dumps(parsed, indent=2, ensure_ascii=False), encoding="utf-8")
            print(f"[OK] {path.name} -> {candidate_id}: {len(parsed['experiences'])} experiences")
        except Exception as exc:
            failures.append({"candidate_id": candidate_id, "source_file": path.name, "error": str(exc), "processing_status": "failed"})
            print(f"[FAILED] {path.name}: {exc}")

    role_similarities = build_role_similarity(results)
    combined = {
        "schema_version": "1.1.0",
        "pipeline": "experience_batch_nlp_relevance_role_similarity",
        "scoring_as_of_date": date.today().isoformat(),
        "input_file_count": len(all_files),
        "resume_count": len(files),
        "success_count": len(results),
        "failure_count": len(failures),
        "candidates": results,
        "failures": failures,
        "relevance_analysis_metadata": {
            "job_profiles_used": job_profiles,
            "scoring_version": "2.1.0-explainable-role-similarity",
            "role_similarity_pair_count": len(role_similarities),
            "skill_evidence_note": "For scoring only, missing skills_used values are augmented with matching terms found in the role title, description, and source text. Original parsed skills_used fields are not overwritten.",
            "score_interpretation": {"low": "below 30", "moderate": "30 to below 65", "high": "65 or above"},
            "decision_support_only": True,
        },
        "role_to_role_similarity": role_similarities,
    }
    output = args.output_dir / "all_candidate_experiences.json"
    output.write_text(json.dumps(combined, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"\nLogical resumes: {len(files)}; successful: {len(results)}; failed: {len(failures)}")
    print(f"Role-to-role comparisons: {len(role_similarities)}")
    print(f"Combined output: {output}")
    return 0 if not failures else 1

if __name__ == "__main__":
    raise SystemExit(main())
